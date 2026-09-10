import os
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"

import glob
import re

import numpy as np
import torch
import torch.nn as nn
import torch.nn.functional as F
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font

# Disable Matplotlib background grids globally.
plt.rcParams["axes.grid"] = False


# =========================================================
# Device
# =========================================================
device = torch.device(
    "cuda" if torch.cuda.is_available() else "cpu"
)

print("Using device:", device)


# =========================================================
# Test settings
# =========================================================
WINDOW_FRAMES = 10
SPAN = WINDOW_FRAMES
MAX_FRAMES = 300

# Training used skip_param = 3, so keep every fourth frame, total 75 frames
TEMPORAL_SKIP_PARAM = 3

# Both models are evaluated on the same 121 x 121 data.
TARGET_SPATIAL_SIZE = (121, 121)
SPATIAL_START_MODE = "even"

# Must match the trained 3D U-Net.
BASE_CHANNELS = 16


# =========================================================
# Noise sensitivity settings
# =========================================================
# Relative AWGN level alpha:
#     noise_std = alpha * std(clean wave field)
#
# Noise is added only to the initial wave/pressure input window.
# Ground truth is never modified.
NOISE_LEVELS = np.concatenate(
    ([0.0], np.linspace(0.001, 0.1, 20))
)

# Ten reproducible noise seeds.
NOISE_SEEDS = list(range(1, 11))


# =========================================================
# General helpers
# =========================================================
def skip_data(data: np.ndarray, param: int) -> np.ndarray:
    if param < 0:
        raise ValueError("param must be non-negative")
    return data[::param + 1]


def add_awgn_to_wave_window(
    clean_window: np.ndarray,
    reference_wave: np.ndarray,
    noise_level: float,
    seed: int,
):
    """
    Add zero-mean white Gaussian noise to the initial wave window.

    noise_std = noise_level * std(reference_wave)
    """
    if noise_level < 0:
        raise ValueError("noise_level must be non-negative")

    signal_std = float(np.std(reference_wave))

    if not np.isfinite(signal_std):
        raise ValueError("Wave-field standard deviation is not finite")

    if signal_std == 0.0:
        signal_std = 1e-12

    noise_std = float(noise_level * signal_std)

    if noise_std == 0.0:
        return clean_window.copy().astype(np.float32), 0.0

    rng = np.random.default_rng(seed)

    noise = rng.normal(
        loc=0.0,
        scale=noise_std,
        size=clean_window.shape,
    ).astype(np.float32)

    noisy_window = clean_window.astype(np.float32) + noise

    return noisy_window.astype(np.float32), noise_std


# =========================================================
# Spatial downsampling
# Same rule as the original evaluation scripts
# =========================================================
def make_downsample_indices(
    n_original,
    n_target,
    start_mode,
):
    if n_target > n_original:
        raise ValueError(
            f"Target size {n_target} cannot be larger "
            f"than original size {n_original}"
        )

    if start_mode == "even":
        start = 0
    elif start_mode == "odd":
        start = 1
    else:
        raise ValueError(
            "start_mode must be 'even' or 'odd'"
        )

    indices = np.linspace(
        start,
        n_original - 1,
        n_target,
    )

    indices = np.round(indices).astype(int)

    indices = np.clip(
        indices,
        0,
        n_original - 1,
    )

    if len(np.unique(indices)) != len(indices):
        raise ValueError(
            "Duplicate indices produced for "
            f"n_original={n_original}, "
            f"n_target={n_target}, "
            f"start_mode={start_mode}"
        )

    return indices


def spatial_downsample_2d(
    arr,
    target_size,
    start_mode,
):
    if arr.ndim != 2:
        raise ValueError(
            f"Expected 2D array, got {arr.shape}"
        )

    target_x, target_y = target_size

    x_indices = make_downsample_indices(
        arr.shape[0],
        target_x,
        start_mode,
    )

    y_indices = make_downsample_indices(
        arr.shape[1],
        target_y,
        start_mode,
    )

    return arr[np.ix_(x_indices, y_indices)]


def spatial_downsample_3d_xyt(
    arr,
    target_size,
    start_mode,
):
    if arr.ndim != 3:
        raise ValueError(
            f"Expected 3D array, got {arr.shape}"
        )

    target_x, target_y = target_size

    x_indices = make_downsample_indices(
        arr.shape[0],
        target_x,
        start_mode,
    )

    y_indices = make_downsample_indices(
        arr.shape[1],
        target_y,
        start_mode,
    )

    time_indices = np.arange(arr.shape[2])

    return arr[
        np.ix_(
            x_indices,
            y_indices,
            time_indices,
        )
    ]


def downsample_spatial_array(
    arr,
    target_size,
    start_mode,
):
    if arr.ndim == 2:
        return spatial_downsample_2d(
            arr,
            target_size,
            start_mode,
        )

    if arr.ndim == 3:
        return spatial_downsample_3d_xyt(
            arr,
            target_size,
            start_mode,
        )

    raise ValueError(
        f"Expected a 2D or 3D array, got {arr.shape}"
    )


# =========================================================
# File discovery
# =========================================================
def get_test_files(data_folder):
    simulation_files = sorted(
        glob.glob(
            os.path.join(
                data_folder,
                "wave_source_*_xyt_kpa.npy",
            )
        )
    )

    if len(simulation_files) == 0:
        raise FileNotFoundError(
            f"No simulation files found in {data_folder}"
        )

    xcoord_file = os.path.join(
        data_folder,
        "X_coordinates_mm.npy",
    )

    ycoord_file = os.path.join(
        data_folder,
        "Y_coordinates_mm.npy",
    )

    tcoord_file = os.path.join(
        data_folder,
        "T_coordinates_us.npy",
    )

    for path in [
        xcoord_file,
        ycoord_file,
        tcoord_file,
    ]:
        if not os.path.exists(path):
            raise FileNotFoundError(
                f"Missing required file: {path}"
            )

    return (
        simulation_files,
        xcoord_file,
        ycoord_file,
        tcoord_file,
    )


# =========================================================
# Source-specific files
# =========================================================
def source_index_from_sim(sim_file):
    basename = os.path.basename(sim_file)

    match = re.search(
        r"wave_source_(\d+)_xyt_kpa\.npy",
        basename,
    )

    if match is None:
        raise ValueError(
            f"Could not extract source index from {basename}"
        )

    return match.group(1)


def pulse_files_for_sim(sim_file):
    source_index = source_index_from_sim(sim_file)
    folder = os.path.dirname(sim_file)

    pulse_x_file = os.path.join(
        folder,
        f"pulse_source_{source_index}_xcoord.npy",
    )

    pulse_y_file = os.path.join(
        folder,
        f"pulse_source_{source_index}_ycoord.npy",
    )

    if not os.path.exists(pulse_x_file):
        raise FileNotFoundError(
            f"Missing pulse x file: {pulse_x_file}"
        )

    if not os.path.exists(pulse_y_file):
        raise FileNotFoundError(
            f"Missing pulse y file: {pulse_y_file}"
        )

    return pulse_x_file, pulse_y_file


def velocity_file_for_sim(sim_file):
    source_index = source_index_from_sim(sim_file)
    folder = os.path.dirname(sim_file)

    source_velocity_file = os.path.join(
        folder,
        f"velocity_map_source_{source_index}.npy",
    )

    global_velocity_file = os.path.join(
        folder,
        "velocity_map.npy",
    )

    if os.path.exists(source_velocity_file):
        return source_velocity_file

    if os.path.exists(global_velocity_file):
        print(
            f"Warning: {source_velocity_file} not found. "
            f"Using {global_velocity_file}"
        )
        return global_velocity_file

    raise FileNotFoundError(
        "Missing velocity map. Expected either:\n"
        f"  {source_velocity_file}\n"
        "or\n"
        f"  {global_velocity_file}"
    )


# =========================================================
# 3D U-Net architecture
# =========================================================
class DoubleConv3D(nn.Module):
    def __init__(
        self,
        in_channels,
        out_channels,
    ):
        super().__init__()

        if out_channels >= 8:
            num_groups = 8
        elif out_channels >= 4:
            num_groups = 4
        else:
            num_groups = 1

        self.block = nn.Sequential(
            nn.Conv3d(
                in_channels,
                out_channels,
                kernel_size=3,
                padding=1,
                bias=False,
            ),
            nn.GroupNorm(
                num_groups=num_groups,
                num_channels=out_channels,
            ),
            nn.GELU(),
            nn.Conv3d(
                out_channels,
                out_channels,
                kernel_size=3,
                padding=1,
                bias=False,
            ),
            nn.GroupNorm(
                num_groups=num_groups,
                num_channels=out_channels,
            ),
            nn.GELU(),
        )

    def forward(self, x):
        return self.block(x)


class UNet3D(nn.Module):
    def __init__(
        self,
        in_channels=7,
        out_channels=1,
        base_channels=16,
    ):
        super().__init__()

        c1 = base_channels
        c2 = base_channels * 2
        c3 = base_channels * 4
        c4 = base_channels * 8

        self.encoder1 = DoubleConv3D(in_channels, c1)
        self.encoder2 = DoubleConv3D(c1, c2)
        self.encoder3 = DoubleConv3D(c2, c3)
        self.bottleneck = DoubleConv3D(c3, c4)

        self.pool = nn.MaxPool3d(
            kernel_size=(1, 2, 2),
            stride=(1, 2, 2),
        )

        self.decoder3 = DoubleConv3D(c4 + c3, c3)
        self.decoder2 = DoubleConv3D(c3 + c2, c2)
        self.decoder1 = DoubleConv3D(c2 + c1, c1)

        self.output_layer = nn.Conv3d(
            c1,
            out_channels,
            kernel_size=1,
        )

    @staticmethod
    def resize_to_skip(x, skip):
        return F.interpolate(
            x,
            size=skip.shape[2:],
            mode="trilinear",
            align_corners=False,
        )

    def forward(self, x):
        enc1 = self.encoder1(x)
        enc2 = self.encoder2(self.pool(enc1))
        enc3 = self.encoder3(self.pool(enc2))
        bottleneck = self.bottleneck(self.pool(enc3))

        dec3 = self.resize_to_skip(bottleneck, enc3)
        dec3 = torch.cat([dec3, enc3], dim=1)
        dec3 = self.decoder3(dec3)

        dec2 = self.resize_to_skip(dec3, enc2)
        dec2 = torch.cat([dec2, enc2], dim=1)
        dec2 = self.decoder2(dec2)

        dec1 = self.resize_to_skip(dec2, enc1)
        dec1 = torch.cat([dec1, enc1], dim=1)
        dec1 = self.decoder1(dec1)

        return self.output_layer(dec1)


def load_unet_model(model_file):
    model = UNet3D(
        in_channels=7,
        out_channels=1,
        base_channels=BASE_CHANNELS,
    ).to(device)

    state = torch.load(
        model_file,
        map_location=device,
        weights_only=False,
    )

    if isinstance(state, dict) and "model_state_dict" in state:
        state = state["model_state_dict"]

    model.load_state_dict(state)
    model.eval()

    print(f"Loaded 3D U-Net model from: {model_file}")

    return model


# =========================================================
# Fourier U-Net + local-integral model
# =========================================================
class SpectralConv3D(nn.Module):
    def __init__(
        self,
        in_channels,
        out_channels,
        modes_t=8,
        modes_x=16,
        modes_y=16,
    ):
        super().__init__()

        self.in_channels = in_channels
        self.out_channels = out_channels
        self.modes_t = modes_t
        self.modes_x = modes_x
        self.modes_y = modes_y

        scale = 1.0 / max(1, in_channels * out_channels)

        shape = (
            in_channels,
            out_channels,
            modes_t,
            modes_x,
            modes_y,
        )

        self.weight_pp = nn.Parameter(
            scale * torch.randn(*shape, dtype=torch.cfloat)
        )
        self.weight_np = nn.Parameter(
            scale * torch.randn(*shape, dtype=torch.cfloat)
        )
        self.weight_pn = nn.Parameter(
            scale * torch.randn(*shape, dtype=torch.cfloat)
        )
        self.weight_nn = nn.Parameter(
            scale * torch.randn(*shape, dtype=torch.cfloat)
        )

    @staticmethod
    def compl_mul3d(x, weight):
        return torch.einsum(
            "bixyz,ioxyz->boxyz",
            x,
            weight,
        )

    def forward(self, x):
        B, _, T, H, W = x.shape

        x_ft = torch.fft.rfftn(
            x,
            dim=(-3, -2, -1),
            norm="ortho",
        )

        out_ft = torch.zeros(
            B,
            self.out_channels,
            T,
            H,
            W // 2 + 1,
            dtype=torch.cfloat,
            device=x.device,
        )

        mt = min(
            self.modes_t,
            T // 2 if T > 1 else 1,
        )
        mx = min(
            self.modes_x,
            H // 2 if H > 1 else 1,
        )
        my = min(
            self.modes_y,
            W // 2 + 1,
        )

        out_ft[:, :, :mt, :mx, :my] = self.compl_mul3d(
            x_ft[:, :, :mt, :mx, :my],
            self.weight_pp[:, :, :mt, :mx, :my],
        )

        out_ft[:, :, -mt:, :mx, :my] = self.compl_mul3d(
            x_ft[:, :, -mt:, :mx, :my],
            self.weight_np[:, :, :mt, :mx, :my],
        )

        out_ft[:, :, :mt, -mx:, :my] = self.compl_mul3d(
            x_ft[:, :, :mt, -mx:, :my],
            self.weight_pn[:, :, :mt, :mx, :my],
        )

        out_ft[:, :, -mt:, -mx:, :my] = self.compl_mul3d(
            x_ft[:, :, -mt:, -mx:, :my],
            self.weight_nn[:, :, :mt, :mx, :my],
        )

        return torch.fft.irfftn(
            out_ft,
            s=(T, H, W),
            dim=(-3, -2, -1),
            norm="ortho",
        )


class FourierBlock3D(nn.Module):
    def __init__(
        self,
        in_channels,
        out_channels,
        modes_t=8,
        modes_x=16,
        modes_y=16,
    ):
        super().__init__()

        self.spectral1 = SpectralConv3D(
            in_channels,
            out_channels,
            modes_t=modes_t,
            modes_x=modes_x,
            modes_y=modes_y,
        )

        self.pointwise1 = nn.Conv3d(
            in_channels,
            out_channels,
            kernel_size=1,
        )

        self.spectral2 = SpectralConv3D(
            out_channels,
            out_channels,
            modes_t=modes_t,
            modes_x=modes_x,
            modes_y=modes_y,
        )

        self.pointwise2 = nn.Conv3d(
            out_channels,
            out_channels,
            kernel_size=1,
        )

        if out_channels >= 8:
            num_groups = 8
        elif out_channels >= 4:
            num_groups = 4
        else:
            num_groups = 1

        self.norm1 = nn.GroupNorm(
            num_groups=num_groups,
            num_channels=out_channels,
        )

        self.norm2 = nn.GroupNorm(
            num_groups=num_groups,
            num_channels=out_channels,
        )

        self.activation = nn.GELU()

    def forward(self, x):
        x = self.activation(
            self.norm1(
                self.spectral1(x) + self.pointwise1(x)
            )
        )

        x = self.activation(
            self.norm2(
                self.spectral2(x) + self.pointwise2(x)
            )
        )

        return x


class LocalIntegralResampler2D(nn.Module):
    def __init__(
        self,
        channels,
        target_size=None,
        radius=0.08,
        num_basis=8,
    ):
        super().__init__()

        self.channels = channels

        if target_size is None:
            self.target_size = None
        elif isinstance(target_size, int):
            self.target_size = (target_size, target_size)
        else:
            self.target_size = tuple(target_size)

        self.radius = float(radius)
        self.num_basis = int(num_basis)

        if self.radius <= 0.0:
            raise ValueError("radius must be > 0")

        if self.num_basis < 2:
            raise ValueError("num_basis must be >= 2")

        centers = torch.linspace(
            0.0,
            1.0,
            self.num_basis,
        )

        init = torch.exp(
            -0.5 * (centers / 0.45) ** 2
        )

        self.theta = nn.Parameter(
            init[None, :].repeat(
                channels,
                1,
            )
        )

    def _basis(self, distance):
        centers = torch.linspace(
            0.0,
            self.radius,
            self.num_basis,
            device=distance.device,
            dtype=distance.dtype,
        )

        width = self.radius / (self.num_basis - 1)

        phi = (
            1.0
            - torch.abs(
                distance[..., None] - centers
            )
            / (width + 1e-12)
        )

        return torch.clamp(phi, min=0.0)

    def _make_axis_weights(
        self,
        n_source,
        n_target,
        device,
        dtype,
    ):
        if n_source < 1 or n_target < 1:
            raise ValueError(
                "Source and target sizes must be positive."
            )

        source_coord = torch.linspace(
            0.0,
            1.0,
            n_source,
            device=device,
            dtype=dtype,
        )

        target_coord = torch.linspace(
            0.0,
            1.0,
            n_target,
            device=device,
            dtype=dtype,
        )

        distance = torch.abs(
            target_coord[:, None]
            - source_coord[None, :]
        )

        phi = self._basis(distance)

        weights = torch.einsum(
            "cl,tsl->cts",
            self.theta.to(dtype=dtype),
            phi,
        )

        support = (
            distance <= self.radius
        ).to(dtype)

        weights = weights * support[None, :, :]

        if n_source > 1:
            q = 1.0 / (n_source - 1)
        else:
            q = 1.0

        weights = weights * q

        return weights

    def forward(
        self,
        x,
        target_size=None,
    ):
        B, C, T, H, W = x.shape

        if C != self.channels:
            raise RuntimeError(
                f"Expected {self.channels} channels, got {C}"
            )

        if target_size is None:
            if self.target_size is None:
                raise ValueError(
                    "target_size must be provided for a "
                    "dynamic LocalIntegralResampler2D."
                )

            H_target, W_target = self.target_size
        else:
            H_target, W_target = target_size

        Wh = self._make_axis_weights(
            H,
            H_target,
            x.device,
            x.dtype,
        )

        Ww = self._make_axis_weights(
            W,
            W_target,
            x.device,
            x.dtype,
        )

        x = torch.einsum(
            "cah,bcthw->bctaw",
            Wh,
            x,
        )

        x = torch.einsum(
            "cdw,bctaw->bctad",
            Ww,
            x,
        )

        return x


class FourierUNet3D(nn.Module):
    def __init__(
        self,
        in_channels=7,
        out_channels=1,
        base_channels=16,
        modes_t=8,
        modes_x=16,
        modes_y=16,
        grid1=60,
        grid2=30,
        grid3=15,
        integral_radius=0.08,
        integral_num_basis=8,
    ):
        super().__init__()

        c1 = base_channels
        c2 = 24
        c3 = 32
        c4 = 48

        self.grid1 = int(grid1)
        self.grid2 = int(grid2)
        self.grid3 = int(grid3)

        self.encoder1 = FourierBlock3D(
            in_channels,
            c1,
            modes_t,
            modes_x,
            modes_y,
        )

        self.encoder2 = FourierBlock3D(
            c1,
            c2,
            modes_t,
            modes_x,
            modes_y,
        )

        self.encoder3 = FourierBlock3D(
            c2,
            c3,
            modes_t,
            modes_x,
            modes_y,
        )

        self.bottleneck = FourierBlock3D(
            c3,
            c4,
            modes_t,
            modes_x,
            modes_y,
        )

        self.down1 = LocalIntegralResampler2D(
            channels=c1,
            target_size=(self.grid1, self.grid1),
            radius=integral_radius,
            num_basis=integral_num_basis,
        )

        self.down2 = LocalIntegralResampler2D(
            channels=c2,
            target_size=(self.grid2, self.grid2),
            radius=integral_radius,
            num_basis=integral_num_basis,
        )

        self.down3 = LocalIntegralResampler2D(
            channels=c3,
            target_size=(self.grid3, self.grid3),
            radius=integral_radius,
            num_basis=integral_num_basis,
        )

        self.decoder3 = FourierBlock3D(
            c4 + c3,
            c3,
            modes_t,
            modes_x,
            modes_y,
        )

        self.decoder2 = FourierBlock3D(
            c3 + c2,
            c2,
            modes_t,
            modes_x,
            modes_y,
        )

        self.decoder1 = FourierBlock3D(
            c2 + c1,
            c1,
            modes_t,
            modes_x,
            modes_y,
        )

        self.up3 = LocalIntegralResampler2D(
            channels=c4,
            target_size=(self.grid2, self.grid2),
            radius=integral_radius,
            num_basis=integral_num_basis,
        )

        self.up2 = LocalIntegralResampler2D(
            channels=c3,
            target_size=(self.grid1, self.grid1),
            radius=integral_radius,
            num_basis=integral_num_basis,
        )

        self.up1 = LocalIntegralResampler2D(
            channels=c2,
            target_size=None,
            radius=integral_radius,
            num_basis=integral_num_basis,
        )

        self.output_layer = nn.Conv3d(
            c1,
            out_channels,
            kernel_size=1,
        )

    def forward(self, x):
        _, _, _, H_native, W_native = x.shape

        enc1 = self.encoder1(x)

        x = self.down1(enc1)
        enc2 = self.encoder2(x)

        x = self.down2(enc2)
        enc3 = self.encoder3(x)

        x = self.down3(enc3)
        bottleneck = self.bottleneck(x)

        dec3 = self.up3(bottleneck)

        if dec3.shape[2:] != enc3.shape[2:]:
            raise RuntimeError(
                "up3 shape does not match skip 3: "
                f"{tuple(dec3.shape)} vs {tuple(enc3.shape)}"
            )

        dec3 = torch.cat([dec3, enc3], dim=1)
        dec3 = self.decoder3(dec3)

        dec2 = self.up2(dec3)

        if dec2.shape[2:] != enc2.shape[2:]:
            raise RuntimeError(
                "up2 shape does not match skip 2: "
                f"{tuple(dec2.shape)} vs {tuple(enc2.shape)}"
            )

        dec2 = torch.cat([dec2, enc2], dim=1)
        dec2 = self.decoder2(dec2)

        dec1 = self.up1(
            dec2,
            target_size=(H_native, W_native),
        )

        if dec1.shape[2:] != enc1.shape[2:]:
            raise RuntimeError(
                "up1 shape does not match skip 1: "
                f"{tuple(dec1.shape)} vs {tuple(enc1.shape)}"
            )

        dec1 = torch.cat([dec1, enc1], dim=1)
        dec1 = self.decoder1(dec1)

        return self.output_layer(dec1)


def load_fourier_unet_model(
    model_file: str,
    config_file: str = None,
):
    config = {
        "in_channels": 7,
        "out_channels": 1,
        "base_channels": 16,
        "modes_t": 24,
        "modes_x": 24,
        "modes_y": 24,
        "window_size": 10,
        "skip_param": 3,
        "loss_crop_size": 121,
        "integral_grid_1": 60,
        "integral_grid_2": 30,
        "integral_grid_3": 15,
        "integral_radius": 0.08,
        "integral_num_basis": 8,
    }

    if config_file is not None and os.path.exists(config_file):
        saved_config = np.load(
            config_file,
            allow_pickle=True,
        ).item()

        for key in config:
            if key in saved_config:
                config[key] = saved_config[key]

        print(
            f"Loaded Fourier U-Net configuration from: {config_file}"
        )
    else:
        print(
            "Fourier U-Net config file not found; using default "
            "architecture settings from the supplied training code."
        )

    model = FourierUNet3D(
        in_channels=int(config["in_channels"]),
        out_channels=int(config["out_channels"]),
        base_channels=int(config["base_channels"]),
        modes_t=int(config["modes_t"]),
        modes_x=int(config["modes_x"]),
        modes_y=int(config["modes_y"]),
        grid1=int(config["integral_grid_1"]),
        grid2=int(config["integral_grid_2"]),
        grid3=int(config["integral_grid_3"]),
        integral_radius=float(config["integral_radius"]),
        integral_num_basis=int(config["integral_num_basis"]),
    ).to(device)

    state = torch.load(
        model_file,
        map_location=device,
        weights_only=False,
    )

    if isinstance(state, dict) and "model_state_dict" in state:
        state = state["model_state_dict"]

    model.load_state_dict(state)
    model.eval()

    print(
        f"Loaded Fourier U-Net + Local Integral model from: {model_file}"
    )

    return model


# =========================================================
# Prepare time coordinate
# =========================================================
def prepare_tcoord_xyt(
    tcoord_raw,
    original_shape,
):
    H, W, T = original_shape

    if tcoord_raw.ndim == 3:
        if tcoord_raw.shape != original_shape:
            raise ValueError(
                f"3D tcoord shape {tcoord_raw.shape} "
                f"does not match simulation {original_shape}"
            )
        return tcoord_raw

    if tcoord_raw.ndim == 2:
        if tcoord_raw.shape != (H, W):
            raise ValueError(
                f"2D tcoord shape {tcoord_raw.shape} "
                f"does not match spatial shape {(H, W)}"
            )

        return np.tile(
            tcoord_raw[:, :, None],
            (1, 1, T),
        )

    raise ValueError(
        f"Invalid tcoord shape: {tcoord_raw.shape}"
    )


# =========================================================
# Shared autoregressive evaluation
# No files or plots are saved here.
# It returns only the average autoregressive MSE.
# =========================================================
def evaluate_autoregressive_mse(
    sim_file,
    xcoord_file,
    ycoord_file,
    tcoord_file,
    model,
    noise_level,
    noise_seed,
):
    sim_name = os.path.splitext(
        os.path.basename(sim_file)
    )[0]

    # -----------------------------------------------------
    # Load full-resolution arrays
    # -----------------------------------------------------
    sim_full_xyt = np.load(sim_file).astype(np.float32)

    if sim_full_xyt.ndim != 3:
        raise ValueError(
            "Expected simulation shape (x,y,t), "
            f"got {sim_full_xyt.shape}"
        )

    original_H, original_W, original_T = sim_full_xyt.shape

    original_T_used = min(
        MAX_FRAMES,
        original_T,
    )

    sim_full_xyt = sim_full_xyt[
        :,
        :,
        :original_T_used,
    ]

    velocity_file = velocity_file_for_sim(sim_file)
    pulse_x_file, pulse_y_file = pulse_files_for_sim(sim_file)

    velocity_full = np.load(velocity_file).astype(np.float32)
    pulse_x_full = np.load(pulse_x_file).astype(np.float32)
    pulse_y_full = np.load(pulse_y_file).astype(np.float32)
    xcoord_full = np.load(xcoord_file).astype(np.float32)
    ycoord_full = np.load(ycoord_file).astype(np.float32)
    tcoord_raw = np.load(tcoord_file).astype(np.float32)

    for name, arr in [
        ("velocity", velocity_full),
        ("pulse_x", pulse_x_full),
        ("pulse_y", pulse_y_full),
        ("xcoord", xcoord_full),
        ("ycoord", ycoord_full),
    ]:
        if arr.shape != (original_H, original_W):
            raise ValueError(
                f"{name} shape {arr.shape} does not "
                f"match simulation spatial shape "
                f"{(original_H, original_W)}"
            )

    tcoord_full_xyt = prepare_tcoord_xyt(
        tcoord_raw=tcoord_raw,
        original_shape=(
            original_H,
            original_W,
            original_T,
        ),
    )

    tcoord_full_xyt = tcoord_full_xyt[
        :,
        :,
        :original_T_used,
    ]

    # -----------------------------------------------------
    # Spatial downsampling to 121 x 121
    # -----------------------------------------------------
    sim_ds_xyt = downsample_spatial_array(
        sim_full_xyt,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    velocity_ds = downsample_spatial_array(
        velocity_full,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    pulse_x_ds = downsample_spatial_array(
        pulse_x_full,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    pulse_y_ds = downsample_spatial_array(
        pulse_y_full,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    xcoord_ds = downsample_spatial_array(
        xcoord_full,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    ycoord_ds = downsample_spatial_array(
        ycoord_full,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    tcoord_ds_xyt = downsample_spatial_array(
        tcoord_full_xyt,
        TARGET_SPATIAL_SIZE,
        SPATIAL_START_MODE,
    ).astype(np.float32)

    # Convert (x,y,t) -> (t,x,y)
    sim = np.transpose(
        sim_ds_xyt,
        (2, 0, 1),
    )

    tcoord = np.transpose(
        tcoord_ds_xyt,
        (2, 0, 1),
    )

    # Same temporal skipping as training.
    sim = skip_data(
        sim,
        TEMPORAL_SKIP_PARAM,
    ).astype(np.float32)

    tcoord = skip_data(
        tcoord,
        TEMPORAL_SKIP_PARAM,
    ).astype(np.float32)

    T, H, W = sim.shape

    if (H, W) != TARGET_SPATIAL_SIZE:
        raise RuntimeError(
            f"Expected spatial size {TARGET_SPATIAL_SIZE}, "
            f"got {(H, W)}"
        )

    if T < 2 * WINDOW_FRAMES:
        raise ValueError(
            f"Only {T} effective frames remain after "
            f"temporal skipping, but at least "
            f"{2 * WINDOW_FRAMES} are needed."
        )

    # -----------------------------------------------------
    # Repeat spatial maps through effective time
    # -----------------------------------------------------
    velocity = np.tile(
        velocity_ds[None, :, :],
        (T, 1, 1),
    )

    pulse_x = np.tile(
        pulse_x_ds[None, :, :],
        (T, 1, 1),
    )

    pulse_y = np.tile(
        pulse_y_ds[None, :, :],
        (T, 1, 1),
    )

    xcoord = np.tile(
        xcoord_ds[None, :, :],
        (T, 1, 1),
    )

    ycoord = np.tile(
        ycoord_ds[None, :, :],
        (T, 1, 1),
    )

    if tcoord.shape != sim.shape:
        raise ValueError(
            f"tcoord shape {tcoord.shape} "
            f"does not match sim shape {sim.shape}"
        )

    def to_5d_tensor(arr):
        return (
            torch.from_numpy(arr)
            .unsqueeze(0)
            .unsqueeze(0)
            .to(device)
        )

    velocity_t = to_5d_tensor(velocity)
    pulse_x_t = to_5d_tensor(pulse_x)
    pulse_y_t = to_5d_tensor(pulse_y)
    xcoord_t = to_5d_tensor(xcoord)
    ycoord_t = to_5d_tensor(ycoord)
    tcoord_t = to_5d_tensor(tcoord)

    # -----------------------------------------------------
    # Add AWGN only to initial wave window
    # -----------------------------------------------------
    clean_initial_wave = sim[:WINDOW_FRAMES]

    noisy_initial_wave, noise_std = add_awgn_to_wave_window(
        clean_window=clean_initial_wave,
        reference_wave=sim,
        noise_level=noise_level,
        seed=noise_seed,
    )

    current_wave = to_5d_tensor(noisy_initial_wave)

    current_input = torch.cat(
        [
            current_wave,
            velocity_t[:, :, :WINDOW_FRAMES],
            pulse_x_t[:, :, :WINDOW_FRAMES],
            pulse_y_t[:, :, :WINDOW_FRAMES],
            xcoord_t[:, :, :WINDOW_FRAMES],
            ycoord_t[:, :, :WINDOW_FRAMES],
            tcoord_t[:, :, :WINDOW_FRAMES],
        ],
        dim=1,
    )

    expected_input_shape = (
        1,
        7,
        WINDOW_FRAMES,
        H,
        W,
    )

    if current_input.shape != expected_input_shape:
        raise RuntimeError(
            f"Expected input shape {expected_input_shape}, "
            f"got {tuple(current_input.shape)}"
        )

    # Ground-truth observed window remains in rollout so MSE is
    # evaluated only for predicted frames below.
    predictions = [sim[:WINDOW_FRAMES]]

    current_start = 0

    # -----------------------------------------------------
    # Autoregressive rollout
    # -----------------------------------------------------
    with torch.inference_mode():
        while (current_start + WINDOW_FRAMES) < T:
            output = model(current_input)

            expected_output_shape = (
                1,
                1,
                WINDOW_FRAMES,
                H,
                W,
            )

            if output.shape != expected_output_shape:
                raise RuntimeError(
                    f"Expected output shape {expected_output_shape}, "
                    f"got {tuple(output.shape)}"
                )

            span_effective = min(
                SPAN,
                T - (current_start + WINDOW_FRAMES),
            )

            next_prediction = output[
                :,
                :,
                :span_effective,
            ]

            predictions.append(
                next_prediction[0, 0]
                .cpu()
                .numpy()
                .astype(np.float32)
            )

            old_wave = current_input[
                :,
                0:1,
                span_effective:,
            ]

            new_wave = torch.cat(
                [old_wave, next_prediction],
                dim=2,
            )

            new_start = current_start + span_effective
            new_end = new_start + WINDOW_FRAMES

            def get_window(tensor):
                window = tensor[
                    :,
                    :,
                    new_start:new_end,
                ]

                if window.shape[2] < WINDOW_FRAMES:
                    missing = WINDOW_FRAMES - window.shape[2]

                    if window.shape[2] == 0:
                        raise RuntimeError(
                            "Cannot pad an empty conditioning window."
                        )

                    padding = window[
                        :,
                        :,
                        -1:,
                    ].repeat(
                        1,
                        1,
                        missing,
                        1,
                        1,
                    )

                    window = torch.cat(
                        [window, padding],
                        dim=2,
                    )

                return window

            current_input = torch.cat(
                [
                    new_wave,
                    get_window(velocity_t),
                    get_window(pulse_x_t),
                    get_window(pulse_y_t),
                    get_window(xcoord_t),
                    get_window(ycoord_t),
                    get_window(tcoord_t),
                ],
                dim=1,
            )

            current_start += span_effective

    rollout = np.concatenate(
        predictions,
        axis=0,
    )[:T]

    if rollout.shape != sim.shape:
        raise RuntimeError(
            f"Rollout shape {rollout.shape} does not "
            f"match ground truth shape {sim.shape}"
        )

    frame_mses = np.mean(
        (
            rollout[WINDOW_FRAMES:]
            - sim[WINDOW_FRAMES:]
        ) ** 2,
        axis=(1, 2),
    )

    average_mse = float(np.mean(frame_mses))

    print(
        f"{sim_name} | alpha={noise_level:.4f} | "
        f"noise_std={noise_std:.6e} | MSE={average_mse:.6e}"
    )

    return average_mse


# =========================================================
# Run sensitivity for one model
# =========================================================
def run_model_sensitivity(
    simulation_files,
    xcoord_file,
    ycoord_file,
    tcoord_file,
    model,
    model_label,
):
    """
    Returns a dictionary containing:
      - per-simulation MSE for every seed/noise level
      - seed-level mean MSE across simulations
      - normalized MSE per seed
      - across-seed mean/std curves

    Normalization is exactly the same convention as the original code:

        normalized_MSE(seed, alpha)
            = mean_MSE(seed, alpha) / mean_MSE(seed, 0)
    """

    noise_array = np.asarray(
        NOISE_LEVELS,
        dtype=np.float64,
    )

    n_seeds = len(NOISE_SEEDS)
    n_noise = len(noise_array)
    n_simulations = len(simulation_files)

    # Shape: (seed, noise, simulation)
    per_sim_mse = np.zeros(
        (n_seeds, n_noise, n_simulations),
        dtype=np.float64,
    )

    simulation_names = [
        os.path.basename(path)
        for path in simulation_files
    ]

    for s_idx, base_seed in enumerate(NOISE_SEEDS):
        print("\n==================================================")
        print(
            f"{model_label} | SEED {base_seed} "
            f"({s_idx + 1}/{n_seeds})"
        )
        print("==================================================")

        for n_idx, noise_level in enumerate(noise_array):
            print(
                f"\n{model_label} | seed={base_seed} | "
                f"alpha={noise_level:.4f} "
                f"({noise_level * 100.0:.2f}% of wave std)"
            )

            for sim_idx, simulation_file in enumerate(simulation_files):
                # This intentionally matches the behavior in your
                # current scripts: the same base seed is used for all
                # noise levels and all simulations for a given seed curve.
                # Both models therefore receive the same AWGN realization
                # for a given seed/noise/simulation input shape.
                experiment_seed = base_seed

                mse = evaluate_autoregressive_mse(
                    sim_file=simulation_file,
                    xcoord_file=xcoord_file,
                    ycoord_file=ycoord_file,
                    tcoord_file=tcoord_file,
                    model=model,
                    noise_level=float(noise_level),
                    noise_seed=int(experiment_seed),
                )

                per_sim_mse[
                    s_idx,
                    n_idx,
                    sim_idx,
                ] = mse

    # Mean across simulations for each seed/noise pair.
    seed_mean_mse = np.mean(
        per_sim_mse,
        axis=2,
    )

    baseline_candidates = np.where(
        np.isclose(noise_array, 0.0)
    )[0]

    if baseline_candidates.size == 0:
        raise RuntimeError(
            "NOISE_LEVELS must contain a clean 0.0 baseline."
        )

    baseline_idx = int(baseline_candidates[0])

    seed_baselines = seed_mean_mse[:, baseline_idx]

    if np.any(seed_baselines <= 0.0):
        raise RuntimeError(
            f"Invalid clean baseline MSE(s): {seed_baselines}"
        )

    # Normalize every seed curve by its own clean baseline.
    seed_normalized_mse = (
        seed_mean_mse
        / seed_baselines[:, None]
    )

    seed_percent_increase = (
        seed_normalized_mse - 1.0
    ) * 100.0

    # Same as original scripts: population std, ddof=0.
    mean_mse = np.mean(seed_mean_mse, axis=0)
    std_mse = np.std(seed_mean_mse, axis=0)

    mean_normalized_mse = np.mean(
        seed_normalized_mse,
        axis=0,
    )

    std_normalized_mse = np.std(
        seed_normalized_mse,
        axis=0,
    )

    mean_percent_increase = np.mean(
        seed_percent_increase,
        axis=0,
    )

    std_percent_increase = np.std(
        seed_percent_increase,
        axis=0,
    )

    return {
        "model_label": model_label,
        "noise_array": noise_array,
        "simulation_names": simulation_names,
        "per_sim_mse": per_sim_mse,
        "seed_mean_mse": seed_mean_mse,
        "seed_normalized_mse": seed_normalized_mse,
        "seed_percent_increase": seed_percent_increase,
        "mean_mse": mean_mse,
        "std_mse": std_mse,
        "mean_normalized_mse": mean_normalized_mse,
        "std_normalized_mse": std_normalized_mse,
        "mean_percent_increase": mean_percent_increase,
        "std_percent_increase": std_percent_increase,
    }


# =========================================================
# Save one Excel workbook per model
# =========================================================
def save_model_results_excel(
    results,
    excel_path,
):
    """
    Saves one Excel workbook for one model with three sheets:

      1. per_simulation
         seed, noise level, simulation, MSE

      2. seed_summary
         one row per seed/noise level after averaging simulations,
         including normalized MSE

      3. aggregate
         across-seed mean/std at every noise level
    """

    wb = Workbook()

    # -----------------------------------------------------
    # Sheet 1: per simulation
    # -----------------------------------------------------
    ws = wb.active
    ws.title = "per_simulation"

    headers = [
        "seed",
        "noise_level",
        "noise_percent",
        "simulation",
        "mse",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    noise_array = results["noise_array"]
    simulation_names = results["simulation_names"]
    per_sim_mse = results["per_sim_mse"]

    for s_idx, seed in enumerate(NOISE_SEEDS):
        for n_idx, noise_level in enumerate(noise_array):
            for sim_idx, simulation_name in enumerate(simulation_names):
                ws.append(
                    [
                        int(seed),
                        float(noise_level),
                        float(noise_level * 100.0),
                        simulation_name,
                        float(per_sim_mse[s_idx, n_idx, sim_idx]),
                    ]
                )

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 34
    ws.column_dimensions["E"].width = 18

    # -----------------------------------------------------
    # Sheet 2: seed-level summary
    # -----------------------------------------------------
    ws = wb.create_sheet("seed_summary")

    headers = [
        "seed",
        "noise_level",
        "noise_percent",
        "mean_mse_across_simulations",
        "normalized_mse",
        "percent_mse_increase",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    seed_mean_mse = results["seed_mean_mse"]
    seed_normalized_mse = results["seed_normalized_mse"]
    seed_percent_increase = results["seed_percent_increase"]

    for s_idx, seed in enumerate(NOISE_SEEDS):
        for n_idx, noise_level in enumerate(noise_array):
            ws.append(
                [
                    int(seed),
                    float(noise_level),
                    float(noise_level * 100.0),
                    float(seed_mean_mse[s_idx, n_idx]),
                    float(seed_normalized_mse[s_idx, n_idx]),
                    float(seed_percent_increase[s_idx, n_idx]),
                ]
            )

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 24

    # -----------------------------------------------------
    # Sheet 3: aggregate across seeds
    # -----------------------------------------------------
    ws = wb.create_sheet("aggregate")

    headers = [
        "noise_level",
        "noise_percent",
        "mean_mse",
        "std_mse",
        "mean_normalized_mse",
        "std_normalized_mse",
        "mean_percent_mse_increase",
        "std_percent_mse_increase",
        "num_seeds",
        "num_simulations",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for i, noise_level in enumerate(noise_array):
        ws.append(
            [
                float(noise_level),
                float(noise_level * 100.0),
                float(results["mean_mse"][i]),
                float(results["std_mse"][i]),
                float(results["mean_normalized_mse"][i]),
                float(results["std_normalized_mse"][i]),
                float(results["mean_percent_increase"][i]),
                float(results["std_percent_increase"][i]),
                len(NOISE_SEEDS),
                len(simulation_names),
            ]
        )

    ws.freeze_panes = "A2"

    for column_letter, width in {
        "A": 15,
        "B": 15,
        "C": 18,
        "D": 18,
        "E": 22,
        "F": 22,
        "G": 28,
        "H": 28,
        "I": 12,
        "J": 16,
    }.items():
        ws.column_dimensions[column_letter].width = width

    wb.save(excel_path)

    print(f"Saved Excel: {excel_path}")


# =========================================================
# Combined normalized-MSE plot
# This is the ONLY plot saved by this script.
# =========================================================
def save_combined_normalized_plot(
    unet_results,
    fourier_results,
    output_path,
):
    noise_array = unet_results["noise_array"]

    if not np.allclose(
        noise_array,
        fourier_results["noise_array"],
    ):
        raise RuntimeError(
            "The two models do not have identical noise levels."
        )

    fig, ax = plt.subplots(figsize=(8, 3.5))

    # 3D U-Net: red mean curve
    unet_line, = ax.plot(
        noise_array,
        unet_results["mean_normalized_mse"],
        marker="o",
        linewidth=2,
        markersize=5,
        color="red",
        label="U-Net",
        zorder=2,
    )

    # 3D U-Net: shaded ±1 std region
    ax.fill_between(
        noise_array,
        unet_results["mean_normalized_mse"]
        - unet_results["std_normalized_mse"],
        unet_results["mean_normalized_mse"]
        + unet_results["std_normalized_mse"],
        alpha=0.15,
        color=unet_line.get_color(),
        zorder=1,
    )

    # Fourier U-Net + local integral: green mean curve
    fourier_line, = ax.plot(
        noise_array,
        fourier_results["mean_normalized_mse"],
        marker="o",
        linewidth=2,
        markersize=5,
        color="green",
        label="Ours",
        zorder=2,
    )

    # Fourier U-Net + local integral: shaded ±1 std region
    ax.fill_between(
        noise_array,
        fourier_results["mean_normalized_mse"]
        - fourier_results["std_normalized_mse"],
        fourier_results["mean_normalized_mse"]
        + fourier_results["std_normalized_mse"],
        alpha=0.15,
        color=fourier_line.get_color(),
        zorder=1,
    )

    # Baseline/reference style
    ax.axhline(
        1.0,
        linestyle="--",
        linewidth=1,
    )

    # y-axis range 
    ax.set_ylim(0.999, 1.003)

    ax.set_xlabel(r"$\alpha$")
    ax.set_ylabel(
        r"Normalized MSE = MSE($\alpha$) / MSE(0)"
    )

    ax.ticklabel_format(
        axis="y",
        style="plain",
        useOffset=False,
    )

    ax.legend()

    # Keep one tick marker at every noise level,
    # but write only 0 on the x axis.
    ax.set_xticks(noise_array)
    ax.set_xticklabels(
        ["0"] + [""] * (len(noise_array) - 1)
    )

    ax.grid(False, which="both", axis="both")

    fig.tight_layout()

    fig.savefig(
        output_path,
        dpi=300,
        bbox_inches="tight",
    )

    plt.close(fig)

    print(f"Saved combined plot: {output_path}")


# =========================================================
# Main
# =========================================================
if __name__ == "__main__":
    # -----------------------------------------------------
    # Input paths
    # -----------------------------------------------------
    data_folder = "Test_data"

    unet_model_path = "unet3d_model.pth"

    fourier_model_path = (
        "fourier_unet3d_local_integral_model.pth"
    )

    fourier_config_path = (
        "fourier_unet3d_local_integral_config.npy"
    )

    # -----------------------------------------------------
    # Output folder
    # Only 3 files are created:
    #   1 combined PNG
    #   2 Excel workbooks, one per model
    # -----------------------------------------------------
    output_folder = "combined_noise_sensitivity_results"
    os.makedirs(output_folder, exist_ok=True)

    unet_excel_path = os.path.join(
        output_folder,
        "unet_noise_sensitivity.xlsx",
    )

    fourier_excel_path = os.path.join(
        output_folder,
        "fourier_unet_local_integral_noise_sensitivity.xlsx",
    )

    combined_plot_path = os.path.join(
        output_folder,
        "combined_normalized_mse_with_std.png",
    )

    # -----------------------------------------------------
    # Test files
    # -----------------------------------------------------
    (
        simulation_files,
        xcoord_file,
        ycoord_file,
        tcoord_file,
    ) = get_test_files(data_folder)

    print("\nSimulation files:")
    for simulation_file in simulation_files:
        print(" ", os.path.basename(simulation_file))

    # -----------------------------------------------------
    # Model 1: 3D U-Net
    # -----------------------------------------------------
    print("\n\n##################################################")
    print("RUNNING 3D U-NET SENSITIVITY")
    print("##################################################")

    unet_model = load_unet_model(
        unet_model_path
    )

    unet_results = run_model_sensitivity(
        simulation_files=simulation_files,
        xcoord_file=xcoord_file,
        ycoord_file=ycoord_file,
        tcoord_file=tcoord_file,
        model=unet_model,
        model_label="U-Net",
    )

    save_model_results_excel(
        unet_results,
        unet_excel_path,
    )

    # Free GPU memory before loading the second model.
    del unet_model
    if torch.cuda.is_available():
        torch.cuda.empty_cache()

    # -----------------------------------------------------
    # Model 2: Fourier U-Net + Local Integral
    # -----------------------------------------------------
    print("\n\n##################################################")
    print("RUNNING FOURIER U-NET + LOCAL INTEGRAL SENSITIVITY")
    print("##################################################")

    fourier_model = load_fourier_unet_model(
        fourier_model_path,
        fourier_config_path,
    )

    fourier_results = run_model_sensitivity(
        simulation_files=simulation_files,
        xcoord_file=xcoord_file,
        ycoord_file=ycoord_file,
        tcoord_file=tcoord_file,
        model=fourier_model,
        model_label="Ours",
    )

    save_model_results_excel(
        fourier_results,
        fourier_excel_path,
    )

    # -----------------------------------------------------
    # One combined normalized-MSE plot with shaded ± std
    # -----------------------------------------------------
    save_combined_normalized_plot(
        unet_results=unet_results,
        fourier_results=fourier_results,
        output_path=combined_plot_path,
    )

    print("\n==================================================")
    print("COMBINED SENSITIVITY ANALYSIS COMPLETE")
    print("==================================================")
    print("Saved only:")
    print(" ", combined_plot_path)
    print(" ", unet_excel_path)
    print(" ", fourier_excel_path)
